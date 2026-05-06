"""Extract PPTX speaker notes via python-pptx (Docling alone misses them)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation


def extract_notes(pptx_path: Path) -> dict[int, str]:
    """Return a mapping ``{slide_index_1based: notes_text}`` for non-empty notes."""
    prs = Presentation(str(pptx_path))
    notes: dict[int, str] = {}
    for idx, slide in enumerate(prs.slides, start=1):
        if not slide.has_notes_slide:
            continue
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            continue
        text = (frame.text or "").strip()
        if text:
            notes[idx] = text
    return notes
