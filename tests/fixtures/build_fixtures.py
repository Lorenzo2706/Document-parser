"""Generate deterministic test fixtures.

Builds:
- ``sample.pptx``: title, bullet list, table, embedded image, speaker notes.
- ``sample.pdf``: multi-slide PDF with bullets and a table.
- ``sample.xlsx``: two worksheets ("Metrics", "Roadmap") with a few rows each.

Run via ``python tests/fixtures/build_fixtures.py`` from the repo root.
"""

from __future__ import annotations

import io
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).parent


def _make_image_with_text(text: str) -> bytes:
    img = Image.new("RGB", (480, 220), color="white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans-Bold.ttf", 28)
    except OSError:
        font = ImageFont.load_default()
    draw.text((20, 90), text, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def build_pptx(target: Path) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]

    # Slide 1 — title + bullets + notes
    s1 = prs.slides.add_slide(blank)
    title = s1.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    title.text_frame.text = "Q1 Roadmap"
    bullets = s1.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(4))
    tf = bullets.text_frame
    tf.text = "Goals for the quarter"
    for line in ("Ship parser MVP", "Onboard 3 design partners", "Cut infra cost 20%"):
        p = tf.add_paragraph()
        p.text = line
        p.level = 1
        p.font.size = Pt(18)
    s1.notes_slide.notes_text_frame.text = (
        "Emphasize that the parser MVP unblocks the design-partner pipeline; "
        "the cost cut depends on the new caching layer landing in week 6."
    )

    # Slide 2 — table + image with text + notes
    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)).text_frame.text = (
        "Metrics"
    )
    tbl = s2.shapes.add_table(rows=3, cols=2, left=Inches(0.5), top=Inches(1.3),
                              width=Inches(5), height=Inches(2)).table
    tbl.cell(0, 0).text = "Metric"
    tbl.cell(0, 1).text = "Value"
    tbl.cell(1, 0).text = "MAU"
    tbl.cell(1, 1).text = "12,400"
    tbl.cell(2, 0).text = "Retention"
    tbl.cell(2, 1).text = "63%"
    img_stream = io.BytesIO(_make_image_with_text("OCR ME PLEASE"))
    s2.shapes.add_picture(img_stream, Inches(6), Inches(1.5), width=Inches(3.5))
    s2.notes_slide.notes_text_frame.text = "Retention is the key narrative this quarter."

    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))


def build_pdf(target: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    target.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(target), pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph("Q1 Roadmap", styles["Title"]))
    story.append(Spacer(1, 18))
    story.append(Paragraph("Goals for the quarter", styles["Heading2"]))
    for line in ("Ship parser MVP", "Onboard 3 design partners", "Cut infra cost 20%"):
        story.append(Paragraph(f"• {line}", styles["BodyText"]))
    story.append(PageBreak())

    story.append(Paragraph("Metrics", styles["Title"]))
    story.append(Spacer(1, 18))
    data = [["Metric", "Value"], ["MAU", "12,400"], ["Retention", "63%"]]
    table = Table(data, colWidths=[180, 180])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("FONTSIZE", (0, 0), (-1, -1), 14),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ]
        )
    )
    story.append(table)

    doc.build(story)


def build_xlsx(target: Path) -> None:
    from openpyxl import Workbook

    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    metrics = wb.active
    metrics.title = "Metrics"
    metrics.append(["Metric", "Value"])
    metrics.append(["MAU", "12,400"])
    metrics.append(["Retention", "63%"])

    roadmap = wb.create_sheet("Roadmap")
    roadmap.append(["Quarter", "Goal"])
    roadmap.append(["Q1", "Ship parser MVP"])
    roadmap.append(["Q2", "Onboard design partners"])

    wb.save(str(target))


if __name__ == "__main__":
    build_pptx(HERE / "sample.pptx")
    build_pdf(HERE / "sample.pdf")
    build_xlsx(HERE / "sample.xlsx")
    print("fixtures written to", HERE)
